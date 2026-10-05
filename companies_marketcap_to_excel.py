#!/usr/bin/env python3
"""CompaniesMarketCapの世界ランキングを3シートのExcelに保存する。

準備（Windowsのコマンドプロンプト）:
    py -m pip install requests beautifulsoup4 openpyxl
実行:
    py companies_marketcap_to_excel.py
保存先指定:
    py companies_marketcap_to_excel.py --output ranking.xlsx

企業名はサイト表記、国名は日本語（未登録の国はサイト表記）。
時価総額はサイトの丸められた表示値を米ドルの数値に変換する。
SSL証明書の検証は有効。取得失敗や表構造の変更時は保存せず終了する。
"""

from __future__ import annotations

import argparse
import re
import sys
import time
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from tempfile import NamedTemporaryFile
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry

URLS = ("https://companiesmarketcap.com/", "https://companiesmarketcap.com/page/2/")
COUNTRIES = {
    "USA": "米国", "United States": "米国", "Japan": "日本",
    "Taiwan": "台湾", "S. Korea": "韓国", "South Korea": "韓国",
    "S. Arabia": "サウジアラビア", "Saudi Arabia": "サウジアラビア",
    "Netherlands": "オランダ", "China": "中国", "UK": "英国",
    "United Kingdom": "英国", "France": "フランス", "Germany": "ドイツ",
    "Switzerland": "スイス", "Canada": "カナダ", "India": "インド",
    "Australia": "オーストラリア", "Ireland": "アイルランド",
    "UAE": "アラブ首長国連邦", "Spain": "スペイン", "Italy": "イタリア",
}


@dataclass(frozen=True)
class Company:
    rank: int
    name: str
    market_cap_usd: Decimal
    country: str


def parse_market_cap(text: str) -> Decimal:
    """例: $3.842 T → 3842000000000。ドル以外の表示は拒否する。"""
    clean = re.sub(r"\s+", "", text).replace(",", "")
    match = re.fullmatch(r"\$(\d+(?:\.\d+)?)([TBMK]?)", clean, re.I)
    if not match:
        raise ValueError(f"時価総額の形式を解釈できません: {text!r}")
    scale = {"T": 10**12, "B": 10**9, "M": 10**6, "K": 10**3, "": 1}
    value = Decimal(match[1]) * scale[match[2].upper()]
    if value <= 0:
        raise ValueError(f"時価総額が正の数ではありません: {text!r}")
    return value


def parse_page(html: str, first_rank: int) -> list[Company]:
    soup = BeautifulSoup(html, "html.parser")
    table = soup.select_one("table.marketcap-table")
    if table is None:
        # CSSクラスが変わっても、企業名の要素を持つ表なら候補にする。
        table = next((t for t in soup.find_all("table") if t.select_one(".company-name")), None)
    if table is None:
        raise ValueError("ランキング表が見つかりません。アクセス制限またはサイト変更の可能性があります。")

    headers = [re.sub(r"\s+", " ", th.get_text(" ", strip=True)).lower()
               for th in table.select("thead th")]

    def column(cells, label):
        if label in headers and headers.index(label) < len(cells):
            return cells[headers.index(label)]
        return None

    records: list[Company] = []
    for tr in table.select("tbody tr"):
        cells = tr.find_all("td", recursive=False)
        if not cells:
            continue
        name_node = tr.select_one(".company-name")
        # 広告・区切り行には企業名がない。100順位の検証で取得漏れを検知する。
        if name_node is None:
            continue
        rank_node = tr.select_one("td.rank-td") or column(cells, "rank")
        if rank_node is None:
            rank_node = next((td for td in cells
                              if re.fullmatch(r"[\d,]+", td.get_text(strip=True))), None)
        if rank_node is None:
            raise ValueError("企業行の順位が見つかりません。")
        rank_text = rank_node.get_text(" ", strip=True).replace(",", "")
        if not re.fullmatch(r"\d+", rank_text):
            raise ValueError(f"順位を解釈できません: {rank_text!r}")
        name = name_node.get_text(" ", strip=True)
        country_node = tr.select_one(".country-name") or column(cells, "country") or cells[-1]
        country = country_node.get_text(" ", strip=True)
        # 国旗絵文字やアイコンを除き、国名のみを残す。
        country = re.sub(r"[^A-Za-z .&()'-]", "", country).strip()
        if not name or not country:
            raise ValueError(f"企業名または国名が空欄です: 順位 {rank_text}")
        cap_node = tr.select_one("td.market-cap-td") or column(cells, "market cap")
        if cap_node is None:
            # 企業名の直後が時価総額。株価のドル表記を誤取得しない。
            name_index = next((i for i, td in enumerate(cells) if td.select_one(".company-name")), None)
            if name_index is not None and name_index + 1 < len(cells):
                cap_node = cells[name_index + 1]
        if cap_node is None:
            raise ValueError(f"時価総額列が見つかりません: 順位 {rank_text}")
        cap = parse_market_cap(cap_node.get_text(" ", strip=True))
        records.append(Company(int(rank_text), name, cap, country))

    expected = set(range(first_rank, first_rank + 100))
    actual = {r.rank for r in records}
    if len(records) != 100 or actual != expected:
        raise ValueError(
            f"順位 {first_rank}～{first_rank + 99} の100件が揃いません。"
            f"取得件数={len(records)}、不足順位={sorted(expected - actual)}"
        )
    return sorted(records, key=lambda r: r.rank)


def fetch_rankings() -> list[Company]:
    retry = Retry(total=3, backoff_factor=1, status_forcelist=[429, 500, 502, 503, 504],
                  allowed_methods=["GET"], respect_retry_after_header=True)
    records: list[Company] = []
    with requests.Session() as session:
        session.headers.update({"User-Agent": "Mozilla/5.0 (compatible; MarketCapExcel/1.0)",
                                "Accept-Language": "en-US,en;q=0.9"})
        session.mount("https://", HTTPAdapter(max_retries=retry))
        for page, url in enumerate(URLS):
            if page:
                time.sleep(2)
            print(f"取得中: {url}")
            response = session.get(url, timeout=(15, 45))
            response.raise_for_status()
            response.encoding = "utf-8"
            records.extend(parse_page(response.text, page * 100 + 1))
    return records


def export_excel(records: list[Company], output: Path, acquired: str) -> list[tuple[str, int]]:
    groups = [
        ("世界1～20位", [r for r in records if 1 <= r.rank <= 20], URLS[0]),
        ("世界100位以内_日本企業", [r for r in records if r.rank <= 100 and r.country == "Japan"], URLS[0]),
        ("世界101～200位_日本企業", [r for r in records if 101 <= r.rank <= 200 and r.country == "Japan"], URLS[1]),
    ]
    if len(groups[0][1]) != 20:
        raise ValueError("世界1～20位の20社が揃っていません。")
    wb = Workbook()
    wb.remove(wb.active)
    counts = []
    for index, (title, companies, url) in enumerate(groups, 1):
        ws = wb.create_sheet(title)
        ws.sheet_view.showGridLines = False
        ws.append([title])
        ws.append([f"取得日時: {acquired}"])
        ws.append([f"出典: {url}"])
        ws["A3"].hyperlink = url
        ws.append(["順位は世界順位。時価総額はサイト表示値を米ドルに換算。"])
        ws.append(["企業名はサイト表記。十億米ドル=10億米ドル。取得時点の値で自動更新はしません。"])
        ws.append([])
        ws.append(["世界順位", "企業名", "時価総額（米ドル）", "時価総額（十億米ドル）", "国名"])
        for item in companies:
            ws.append([item.rank, item.name, float(item.market_cap_usd),
                       float(item.market_cap_usd / Decimal(10**9)),
                       COUNTRIES.get(item.country, item.country)])
        if companies:
            table = Table(displayName=f"Companies{index}", ref=f"A7:E{ws.max_row}")
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            ws.add_table(table)
        else:
            ws["A8"] = "該当企業なし"
        for cell in ws[7]:
            cell.font = Font(name="Meiryo", bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="17365D")
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for row in ws.iter_rows(min_row=8):
            for cell in row:
                cell.font = Font(name="Meiryo", size=11)
                cell.alignment = Alignment(vertical="center")
            row[2].number_format = '"$"#,##0'
            row[3].number_format = '#,##0.00'
            ws.row_dimensions[row[0].row].height = 24
        ws["A1"].font = Font(name="Meiryo", size=16, bold=True, color="17365D")
        for col, width in {"A": 13, "B": 55, "C": 27, "D": 30, "E": 24}.items():
            ws.column_dimensions[col].width = width
        ws.freeze_panes = "C8"
        ws.row_dimensions[7].height = 26
        counts.append((title, len(companies)))

    output = output.expanduser().resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    # 保存途中のエラーで既存ファイルを壊さない。
    with NamedTemporaryFile(dir=output.parent, suffix=".xlsx", delete=False) as temp:
        temporary = Path(temp.name)
    try:
        wb.save(temporary)
        temporary.replace(output)
    finally:
        temporary.unlink(missing_ok=True)
    return counts


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--output", type=Path, help="出力先Excelファイル（省略時はプログラムと同じフォルダ）")
    args = parser.parse_args()
    try:
        try:
            now = datetime.now(ZoneInfo("Asia/Tokyo"))
        except ZoneInfoNotFoundError:
            # WindowsでタイムゾーンDBがない場合も日本時間にする。
            from datetime import timedelta, timezone
            now = datetime.now(timezone(timedelta(hours=9)))
        output = args.output or Path(__file__).resolve().parent / f"企業時価総額_{now:%Y%m%d_%H%M%S}.xlsx"
        if output.suffix.lower() != ".xlsx":
            raise ValueError("出力ファイルの拡張子は .xlsx にしてください。")
        records = fetch_rankings()
        counts = export_excel(records, output, now.strftime("%Y-%m-%d %H:%M:%S JST（取得開始）"))
        for title, count in counts:
            print(f"{title}: {count}社")
        print(f"保存完了: {output.resolve()}")
        return 0
    except PermissionError:
        print("保存できません。出力先のExcelを閉じ、書き込み可能なフォルダで再実行してください。", file=sys.stderr)
    except (requests.RequestException, ValueError, OSError) as exc:
        print(f"処理に失敗しました: {exc}", file=sys.stderr)
        print("アクセス制限やサイト変更の場合は、時間を置くかHTML構造を確認してください。", file=sys.stderr)
    return 1


if __name__ == "__main__":
    raise SystemExit(main())
